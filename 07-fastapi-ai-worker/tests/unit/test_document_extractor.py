"""Unit tests for DocumentExtractor service (PDF, DOCX, TXT, magic bytes, size limits)."""

import io
import pytest
from docx import Document as DocxDocument
import PyPDF2

from app.services.document_extractor import DocumentExtractor


def test_document_extractor_plain_text():
    """Extract plain text successfully."""
    extractor = DocumentExtractor()
    sample_text = "Software Engineer with 5 years experience in Python and FastAPI."
    schema = extractor.extract_from_bytes("resume.txt", sample_text.encode("utf-8"))

    assert schema.extracted_text == sample_text
    assert schema.schema_version == "1.0"
    assert schema.overall_confidence == 100.0


def test_document_extractor_docx_extraction():
    """Extract text from a generated in-memory DOCX file."""
    extractor = DocumentExtractor()
    doc = DocxDocument()
    doc.add_heading("Alice Wonderland - Resume", 0)
    doc.add_paragraph("Skills: Python, TypeScript, PostgreSQL")

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    schema = extractor.extract_from_bytes("alice.docx", docx_bytes)
    assert "Alice Wonderland" in schema.extracted_text
    assert "TypeScript" in schema.extracted_text


def test_document_extractor_pdf_magic_bytes_check():
    """Reject file with .pdf extension but invalid magic bytes."""
    extractor = DocumentExtractor()
    invalid_pdf_content = b"This is not a real PDF file header."

    with pytest.raises(ValueError, match="magic bytes do not match PDF format"):
        extractor.extract_from_bytes("fake.pdf", invalid_pdf_content)


def test_document_extractor_docx_magic_bytes_check():
    """Reject file with .docx extension but invalid PK magic bytes."""
    extractor = DocumentExtractor()
    invalid_docx_content = b"Not a zip or docx file."

    with pytest.raises(ValueError, match="magic bytes do not match ZIP archive"):
        extractor.extract_from_bytes("fake.docx", invalid_docx_content)


def test_document_extractor_oversized_document():
    """Reject document exceeding maximum allowed size."""
    extractor = DocumentExtractor(max_document_size_bytes=100)
    large_content = b"A" * 200

    with pytest.raises(ValueError, match="exceeds max size"):
        extractor.extract_from_bytes("large.txt", large_content)


def test_document_extractor_text_truncation():
    """Text exceeding max_extracted_text_length is safely truncated."""
    extractor = DocumentExtractor(max_extracted_text_length=50)
    long_text = "Python Developer " * 10
    schema = extractor.extract_from_bytes("long.txt", long_text.encode("utf-8"))

    assert len(schema.extracted_text) == 50
